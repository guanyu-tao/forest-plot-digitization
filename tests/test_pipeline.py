"""Small synthetic cases for data-preserving pipeline invariants."""

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import forest_plot as fp
import export_xlsx as ex
from openpyxl import load_workbook


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.work = self.root / "work"
        self.work.mkdir()
        self.output = self.root / "output"
        self.output.mkdir()
        self.pdf = self.root / "synthetic.pdf"
        self.pdf.write_bytes(b"synthetic fixture; never a real scan")
        fp.jwrite(self.work / "inventory.json", {"pdfs": [{"name": self.pdf.name, "stem": self.pdf.stem, "path": str(self.pdf), "pages": 1, "sha256": fp.digest(self.pdf)}]})
        fp.jwrite(self.work / "page_reviews.json", {self.pdf.name: {"1": {"expected_rows": 2, "checked": True, "xy_unit": "cm", "unit_evidence": "synthetic page heading"}}})
        fp.jwrite(fp.ocr_path(self.work, self.pdf.stem, 1), {"tokens": [{"text": "=1+1", "score": 0.8, "poly": [[0, 0], [1, 0], [1, 1], [0, 1]]}]})

    def record(self, row, tag, x, species="尖连蕊茶"):
        return {"pdf": self.pdf.name, "page": 1, "row": row, "checked": True, "sample": "0304", "small": 2, "tag": tag, "species": species, "x": x, "y": 300, "dbh": 2.5, "d1": "1.2g", "d2": None, "d3": None, "d4": None, "d5": None, "d6": None, "remark": "枯立；已采", "date": "2026-09-01"}

    def test_finalize_export_validate_and_literal_ocr(self):
        rows = [self.record(1, "HHBX03040001", 250), self.record(2, "HHBX03040001", 600)]
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))
        payload = fp.jread(self.work / "final_payload.json")
        output = payload["outputs"][0]
        self.assertEqual(output["record_count"], 2)
        self.assertEqual(output["pending_count"], 2)  # duplicate tag, plus x out of range
        self.assertEqual(output["main"][0][7], "尖连蕊茶")  # alias is only a suggestion
        self.assertEqual(output["main"][0][8], 2.5)
        self.assertEqual(output["main"][1][8], 6.0)
        self.assertEqual(output["main"][0][12], "g")  # no invented suffix number
        self.assertEqual(output["main"][0][23], "枯立")
        ex.export_one(payload, output, self.output, False, 1)
        fp.validate(SimpleNamespace(workdir=str(self.work), output_dir=str(self.output)))
        book = load_workbook(self.output / output["file_name"])
        self.assertEqual(book.sheetnames, ["主表", "原始OCR", "质检汇总", "修订记录", "物种名录核对"])
        self.assertEqual(book["原始OCR"]["D2"].value, "=1+1")
        self.assertEqual(book["原始OCR"]["D2"].data_type, "s")
        with self.assertRaises(FileExistsError):
            ex.export_one(payload, output, self.output, False, 1)

    def test_out_of_range_page_rejected(self):
        rows = [self.record(1, "HHBX03040001", 250), self.record(2, "HHBX03040002", 300)]
        rows[1]["page"] = 2
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        with self.assertRaisesRegex(ValueError, "页码范围外"):
            fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))

    def test_unknown_unit_is_not_guessed(self):
        reviews = fp.jread(self.work / "page_reviews.json")
        reviews[self.pdf.name]["1"]["xy_unit"] = "unknown"
        fp.jwrite(self.work / "page_reviews.json", reviews)
        rows = [self.record(1, "HHBX03040001", 250, "已核物种"), self.record(2, "HHBX03040002", 300, "已核物种")]
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))
        main = fp.jread(self.work / "final_payload.json")["outputs"][0]["main"]
        self.assertIsNone(main[0][8])
        self.assertIsNone(main[0][9])
        self.assertIn("单位未由原图确认", main[0][-1])

    def test_unreviewed_record_rejected(self):
        rows = [self.record(1, "HHBX03040001", 250), self.record(2, "HHBX03040002", 300)]
        rows[1]["checked"] = False
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        with self.assertRaisesRegex(ValueError, "尚未图像复核"):
            fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))

    def test_empty_page_exports_and_validates(self):
        reviews = fp.jread(self.work / "page_reviews.json")
        reviews[self.pdf.name]["1"]["expected_rows"] = 0
        fp.jwrite(self.work / "page_reviews.json", reviews)
        fp.write_jsonl(self.work / "reviewed.jsonl", [])
        fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))
        payload = fp.jread(self.work / "final_payload.json")
        ex.export_one(payload, payload["outputs"][0], self.output, False, 1)
        fp.validate(SimpleNamespace(workdir=str(self.work), output_dir=str(self.output)))

    def test_catalog_checks_names_without_auto_correction(self):
        catalog = self.root / "catalog.csv"
        catalog.write_text("vernacularName,scientificName,isAcceptedName,scientificNameID\n银杏,Ginkgo biloba,1,taxon-1\n", encoding="utf-8")
        rows = [self.record(1, "HHBX03040001", 250, "银杏"), self.record(2, "HHBX03040002", 300, "银杏错")]
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        fp.names(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE), catalog=str(catalog), catalog_source="synthetic catalog", catalog_version="test-1"))
        audit = fp.jread(self.work / "species_audit.json")
        by_name = {row["name"]: row for row in audit["audit"]}
        self.assertEqual(audit["catalog"]["sha256"], fp.digest(catalog))
        self.assertEqual(by_name["银杏"]["status"], "精确命中")
        self.assertEqual(by_name["银杏错"]["status"], "名录未命中")
        self.assertIn("银杏", by_name["银杏错"]["candidates"])
        fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))
        main = fp.jread(self.work / "final_payload.json")["outputs"][0]["main"]
        self.assertEqual(main[0][7], "银杏")
        self.assertEqual(main[0][-2], "图像复核通过")
        self.assertEqual(main[1][7], "银杏错")
        self.assertEqual(main[1][-2], "待人工核对")
        self.assertIn("名录未命中", main[1][-1])
        payload = fp.jread(self.work / "final_payload.json")
        ex.export_one(payload, payload["outputs"][0], self.output, False, 1)
        fp.validate(SimpleNamespace(workdir=str(self.work), output_dir=str(self.output)))
        book = load_workbook(self.output / payload["outputs"][0]["file_name"])
        self.assertEqual(book["物种名录核对"]["C2"].value, "精确命中")

    def test_stale_catalog_audit_rejected(self):
        catalog = self.root / "catalog.csv"
        catalog.write_text("中文名,学名\n银杏,Ginkgo biloba\n", encoding="utf-8")
        rows = [self.record(1, "HHBX03040001", 250, "银杏"), self.record(2, "HHBX03040002", 300, "银杏")]
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        fp.names(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE), catalog=str(catalog), catalog_source="synthetic catalog", catalog_version="test-1"))
        rows[1]["species"] = "银杏错"
        fp.write_jsonl(self.work / "reviewed.jsonl", rows)
        with self.assertRaisesRegex(ValueError, "已过期"):
            fp.finalize(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE)))

    def test_catalog_distinguishes_synonym_and_ambiguous_name(self):
        catalog = self.root / "catalog.csv"
        catalog.write_text("vernacularName,scientificName,isAcceptedName,acceptedNameUsageID\n旧名,Species alba,0,id-1\n重名,Species beta,1,id-2\n重名,Species gamma,1,id-3\n", encoding="utf-8")
        fp.write_jsonl(self.work / "reviewed.jsonl", [self.record(1, "HHBX03040001", 250, "旧名"), self.record(2, "HHBX03040002", 300, "重名")])
        fp.names(SimpleNamespace(workdir=str(self.work), profile=str(fp.DEFAULT_PROFILE), catalog=str(catalog), catalog_source="synthetic catalog", catalog_version="test-1"))
        by_name = {row["name"]: row for row in fp.jread(self.work / "species_audit.json")["audit"]}
        self.assertEqual(by_name["旧名"]["status"], "精确命中异名，待核")
        self.assertEqual(by_name["重名"]["status"], "同名对应多个分类单元，待核")


if __name__ == "__main__":
    unittest.main()
