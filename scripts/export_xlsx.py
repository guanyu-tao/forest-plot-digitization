"""Export reviewed forest-plot payloads with the public openpyxl package."""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = {
    "main": ["来源文件", "PDF页码", "页内行号", "调查日期", "样方编号", "小方", "标牌号", "物种名", "X（m）", "Y（m）", "DBH", "D1", "D1分枝类型", "D2", "D2分枝类型", "D3", "D3分枝类型", "D4", "D4分枝类型", "D5", "D5分枝类型", "D6", "D6分枝类型", "备注", "复核状态", "疑点说明"],
    "raw": ["来源文件", "PDF页码", "OCR序号", "原始文本", "OCR置信度", "检测框"],
    "summary": ["来源文件", "PDF页码", "应有条数", "录入待核条数", "X/Y原始单位", "单位依据", "页备注"],
    "changes": ["来源文件", "PDF页码", "页内行号", "标牌号", "字段", "原值", "修订值", "处理", "依据", "置信度"],
    "species_catalog": ["最终中文名", "记录数", "名录核对状态", "名录学名", "近似名称候选", "名录来源", "名录版本"],
}


def write_rows(sheet, rows, width):
    for row in rows:
        if len(row) != width:
            raise ValueError(f"{sheet.title}: expected {width} columns, received {len(row)}")
        sheet.append(row)
        # OCR is untrusted input. Keep leading '=' as literal text, never a formula.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"


def header(sheet, row=1):
    for cell in sheet[row]:
        cell.fill = PatternFill("solid", fgColor="174A67")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28
    sheet.freeze_panes = "A2" if row == 1 else f"A{row + 1}"
    sheet.sheet_view.showGridLines = False


def widths(sheet, widths_by_column):
    for column, width in widths_by_column.items():
        sheet.column_dimensions[column].width = width


def export_one(payload, source, output_dir: Path, overwrite: bool, index: int):
    file_name = source["file_name"]
    if Path(file_name).name != file_name or not file_name.lower().endswith(".xlsx"):
        raise ValueError(f"unsafe Excel filename: {file_name}")
    target = output_dir / file_name
    if target.exists() and not overwrite:
        raise FileExistsError(target)
    wb = Workbook()
    main = wb.active
    main.title = "主表"
    write_rows(main, [HEADERS["main"], *source["main"]], 26)
    header(main)
    widths(main, {"A": 32, "B": 11, "C": 11, "D": 18, "E": 11, "F": 11, "G": 22, "H": 19, "X": 20, "Y": 15, "Z": 65})
    for column in range(9, 24):
        main.column_dimensions[get_column_letter(column)].width = 12
    for row in main.iter_rows(min_row=2, min_col=9, max_col=23):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00##"
    if source["main"]:
        table = Table(displayName=f"ForestRecords{index}", ref=f"A1:Z{main.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        main.add_table(table)
    else:
        main.auto_filter.ref = "A1:Z1"

    raw = wb.create_sheet("原始OCR")
    write_rows(raw, [HEADERS["raw"], *source["raw"]], 6)
    header(raw)
    widths(raw, {"A": 32, "B": 11, "C": 11, "D": 50, "E": 17, "F": 75})

    qa = wb.create_sheet("质检汇总")
    write_rows(qa, [["项目", "结果"], ["来源PDF", Path(source["source_pdf"]).name], ["PDF页数", source["page_count"]], ["录入总条数", source["record_count"]], ["待人工核对条数", source["pending_count"]], ["图像复核通过条数", source["record_count"] - source["pending_count"]], ["表格规则", payload["profile"]]], 2)
    header(qa)
    qa.append([])
    write_rows(qa, [HEADERS["summary"], *source["summary"]], 7)
    header(qa, 9)
    widths(qa, {"A": 32, "B": 20, "C": 20, "D": 20, "E": 20, "F": 56, "G": 56})

    changes = wb.create_sheet("修订记录")
    write_rows(changes, [HEADERS["changes"], *source["changes"]], 10)
    header(changes)
    widths(changes, {"A": 32, "B": 18, "C": 18, "D": 18, "E": 18, "F": 22, "G": 22, "H": 18, "I": 70, "J": 16})

    catalog = wb.create_sheet("物种名录核对")
    write_rows(catalog, [HEADERS["species_catalog"], *source["species_catalog"]], 7)
    header(catalog)
    widths(catalog, {"A": 22, "B": 12, "C": 30, "D": 42, "E": 38, "F": 34, "G": 18})

    temporary = target.with_name(target.name + ".partial")
    if temporary.exists():
        raise FileExistsError(temporary)
    try:
        wb.save(temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return {"output": str(target), "pages": source["page_count"], "rows": source["record_count"], "pending": source["pending_count"]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--payload", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true", help="replace existing Excel outputs")
    args = parser.parse_args()
    payload = json.loads(args.payload.read_text(encoding="utf-8"))
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    for index, source in enumerate(payload["outputs"], 1):
        print(json.dumps(export_one(payload, source, output_dir, args.overwrite, index), ensure_ascii=False))


if __name__ == "__main__":
    main()
