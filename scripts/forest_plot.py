"""Reusable, source-preserving forest-plot scan/OCR/review/QC pipeline.

OCR and grid assignment produce candidates. Finalization requires explicit
page-level and row-level image review records; the script never claims that
OCR alone verified handwriting.
"""

from __future__ import annotations

import argparse
import bisect
import collections
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path


SKILL = Path(__file__).resolve().parents[1]
DEFAULT_PROFILE = SKILL / "references" / "houhe-hhbx.json"
REQUIRED = ("pdf", "page", "row", "checked", "sample", "small", "tag", "species", "x", "y", "dbh", "remark")
NUMBER = re.compile(r"^[+-]?\d+(?:\.\d+)?$")
BRANCH = re.compile(r"^([+-]?\d+(?:\.\d+)?)\s*([gzGZ])?\s*([1-6])?$")
BRANCH_PREFIX = re.compile(r"^([gzGZ])\s*[:：]\s*([+-]?\d+(?:\.\d+)?)$")


def jread(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def jwrite(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def jsonl(path: Path):
    if not path.exists():
        raise FileNotFoundError(path)
    with path.open(encoding="utf-8") as f:
        for n, line in enumerate(f, 1):
            if line.strip():
                try:
                    yield json.loads(line)
                except json.JSONDecodeError as e:
                    raise ValueError(f"{path}:{n}: {e}") from e


def write_jsonl(path: Path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def digest(path: Path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def manifest(workdir: Path):
    doc = jread(workdir / "inventory.json")
    for item in doc["pdfs"]:
        src = Path(item["path"])
        if not src.is_file() or digest(src) != item["sha256"]:
            raise RuntimeError(f"源 PDF 已变化或不可用：{src}")
    return doc


def page_path(workdir: Path, stem: str, page: int):
    return workdir / "pages" / stem / f"p-{page:03}.png"


def ocr_path(workdir: Path, stem: str, page: int):
    return workdir / "ocr" / stem / f"p-{page:03}.json"


def selected_pages(item, limit):
    return range(1, min(item["pages"], limit or item["pages"]) + 1)


def doctor(args):
    ocr_python = args.ocr_python or os.environ.get("PADDLE_OCR_PYTHON") or sys.executable
    result = {
        "python": sys.executable,
        "pdftoppm": shutil.which("pdftoppm"),
        "ocr_python": ocr_python,
        "profile": str(DEFAULT_PROFILE),
    }
    if args.input_dir:
        source = Path(args.input_dir)
        result["input_dir"] = str(source.resolve())
        result["pdf_count"] = len(list(source.glob("*.pdf"))) if source.is_dir() else None
    try:
        probe = subprocess.run([ocr_python, "-c", "import paddle,paddleocr; print(paddle.__version__,paddleocr.__version__)"], capture_output=True, text=True, timeout=75)
        result["paddle_probe"] = probe.stdout.strip() if probe.returncode == 0 else "unavailable; install PaddleOCR in the OCR interpreter"
    except (OSError, subprocess.TimeoutExpired) as exc:
        result["paddle_probe"] = f"unavailable: {type(exc).__name__}"
    print(json.dumps(result, ensure_ascii=False, indent=2))


def inventory(args):
    from pypdf import PdfReader

    source = Path(args.input_dir).resolve()
    workdir = Path(args.workdir).resolve()
    if not source.is_dir():
        raise NotADirectoryError(source)
    if source == workdir or source in workdir.parents or workdir in source.parents:
        raise ValueError("工作目录必须与输入目录分开，避免中间文件混入扫描件目录")
    excludes = set(args.exclude or [])
    items = []
    for pdf in sorted(source.glob("*.pdf"), key=lambda p: p.name.lower()):
        if pdf.name in excludes:
            continue
        pages = len(PdfReader(str(pdf)).pages)
        items.append({"name": pdf.name, "stem": pdf.stem, "path": str(pdf), "pages": pages, "sha256": digest(pdf)})
    if not items:
        raise ValueError("没有可处理的 PDF")
    old = workdir / "inventory.json"
    if old.exists():
        previous = jread(old)
        if previous["pdfs"] != items:
            raise RuntimeError("工作目录已有不同批次的 inventory；请换一个工作目录")
    workdir.mkdir(parents=True, exist_ok=True)
    jwrite(old, {"input_dir": str(source), "pdfs": items, "excluded": sorted(excludes)})
    print(json.dumps({"workdir": str(workdir), "pdfs": [{"name": x["name"], "pages": x["pages"]} for x in items]}, ensure_ascii=False))


def render(args):
    workdir = Path(args.workdir).resolve()
    doc = manifest(workdir)
    exe = shutil.which("pdftoppm")
    if not exe:
        raise RuntimeError("pdftoppm 不在 PATH 中")
    rotations_file = workdir / "page_rotations.json"
    rotations = jread(rotations_file) if rotations_file.exists() else {}
    made = skipped = 0
    for item in doc["pdfs"]:
        for page in selected_pages(item, args.limit_pages):
            target = page_path(workdir, item["stem"], page)
            rotation = rotations.get(item["name"], {}).get(str(page), 0)
            if rotation not in (0, 90, 180, 270):
                raise ValueError(f"旋转角度必须为0/90/180/270：{item['name']} 第{page}页")
            settings = {"source_sha256": item["sha256"], "dpi": args.dpi, "clockwise_rotation": rotation}
            metadata = target.with_suffix(".meta.json")
            if target.exists() and target.stat().st_size > 10000 and not args.force:
                if metadata.exists() and jread(metadata) == settings:
                    skipped += 1
                    continue
                raise RuntimeError(f"已有图片的渲染设置不同：{target}；核对后使用 --force 重渲染")
            target.parent.mkdir(parents=True, exist_ok=True)
            prefix = str(target.with_suffix(""))
            cmd = [exe, "-f", str(page), "-l", str(page), "-r", str(args.dpi), "-png", "-singlefile", item["path"], prefix]
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL)
            if not target.exists() or target.stat().st_size < 10000:
                raise RuntimeError(f"渲染失败：{item['name']} 第{page}页")
            if rotation:
                from PIL import Image
                with Image.open(target) as image:
                    image.rotate(-rotation, expand=True).save(target)
            jwrite(metadata, settings)
            for stale in (ocr_path(workdir, item["stem"], page), ocr_path(workdir, item["stem"], page).with_name(f"p-{page:03}_raw.json")):
                stale.unlink(missing_ok=True)
            made += 1
            print(f"render {item['name']} {page}/{item['pages']}", flush=True)
    print(json.dumps({"rendered": made, "skipped": skipped}, ensure_ascii=False))


def normalized_result(data):
    if isinstance(data, str):
        data = json.loads(data)
    if "res" in data and isinstance(data["res"], dict):
        data = data["res"]
    if "overall_ocr_res" in data:
        data = data["overall_ocr_res"]
    texts = data.get("rec_texts") or []
    scores = data.get("rec_scores") or []
    polys = data.get("dt_polys") or data.get("rec_polys") or []
    tokens = []
    for i, text in enumerate(texts):
        poly = polys[i] if i < len(polys) else None
        if hasattr(poly, "tolist"):
            poly = poly.tolist()
        tokens.append({"text": str(text), "score": float(scores[i]) if i < len(scores) else None, "poly": poly})
    return tokens


def ocr(args):
    workdir = Path(args.workdir).resolve()
    doc = manifest(workdir)
    os.environ.setdefault("PADDLE_PDX_MODEL_SOURCE", "BOS")
    if args.structure:
        from paddleocr import PPStructureV3
        model = PPStructureV3(text_detection_model_name="PP-OCRv6_medium_det", text_recognition_model_name="PP-OCRv6_medium_rec", use_doc_orientation_classify=True, use_doc_unwarping=True, use_textline_orientation=False, use_formula_recognition=False, engine="paddle")
    else:
        from paddleocr import PaddleOCR
        model = PaddleOCR(text_detection_model_name="PP-OCRv6_medium_det", text_recognition_model_name="PP-OCRv6_medium_rec", use_doc_orientation_classify=False, use_doc_unwarping=False, use_textline_orientation=False, engine="paddle")
    made = skipped = 0
    for item in doc["pdfs"]:
        for page in selected_pages(item, args.limit_pages):
            image = page_path(workdir, item["stem"], page)
            target = ocr_path(workdir, item["stem"], page)
            if not image.is_file():
                raise FileNotFoundError(image)
            if target.exists():
                skipped += 1
                continue
            result = next(iter(model.predict(input=str(image))))
            raw = result.json
            if isinstance(raw, str):
                raw = json.loads(raw)
            target.parent.mkdir(parents=True, exist_ok=True)
            jwrite(target, {"pdf": item["name"], "page": page, "model": "PP-StructureV3" if args.structure else "PP-OCRv6", "tokens": normalized_result(raw)})
            jwrite(target.with_name(target.stem + "_raw.json"), raw)
            made += 1
            print(f"ocr {item['name']} {page}/{item['pages']} tokens={len(normalized_result(raw))}", flush=True)
    print(json.dumps({"ocr_pages": made, "skipped": skipped}, ensure_ascii=False))


def clusters(indices, max_gap=2):
    groups = []
    for x in indices:
        if groups and x - groups[-1][-1] <= max_gap:
            groups[-1].append(x)
        else:
            groups.append([x])
    return [round(sum(g) / len(g)) for g in groups]


def grid_lines(image: Path, profile=None):
    import numpy as np
    from PIL import Image

    im = np.asarray(Image.open(image).convert("L"))
    h, w = im.shape
    if profile and profile.get("printed_rows") and profile.get("reference_x_boundaries"):
        try:
            import cv2
            binary = cv2.threshold(im, 190, 255, cv2.THRESH_BINARY_INV)[1]
            detected = cv2.HoughLinesP(binary, 1, np.pi / 180, threshold=max(45, round(90 * w / 1406)), minLineLength=round(min(w, h) * .25), maxLineGap=max(15, round(35 * w / 1406)))
            xhit, yhit = [], []
            if detected is not None:
                for x1, y1, x2, y2 in detected[:, 0]:
                    dx, dy = x2 - x1, y2 - y1
                    if abs(dy) > h * .25 and abs(dx) < abs(dy) * .08:
                        xhit.append(round(x1 + dx * ((h * .5 - y1) / dy)))
                    if abs(dx) > w * .25 and abs(dy) < abs(dx) * .08:
                        yhit.append(round(y1 + dy * ((w * .5 - x1) / dx)))
            x_candidates = clusters(sorted(xhit), max(3, round(8 * w / 1406)))
            y_candidates = clusters(sorted(yhit), max(2, round(5 * h / 1988)))
            template = profile["reference_x_boundaries"]
            best_x = None
            tolerance_x = max(10, round(17 * w / 1406))
            for count in profile.get("allowed_column_boundaries", [len(template)]):
                if count > len(template):
                    continue
                reference = template[:count]
                for left in x_candidates:
                    if not .055 * w <= left <= .15 * w:
                        continue
                    for right in x_candidates:
                        if not .84 * w <= right <= .98 * w or right <= left:
                            continue
                        if (right - left) / w < (.75 if count == 14 else .70):
                            continue
                        predictions = [left + (right - left) * (x - reference[0]) / (reference[-1] - reference[0]) for x in reference]
                        matches = [min(x_candidates, key=lambda v: abs(v - p)) for p in predictions]
                        errors = [abs(x - p) for x, p in zip(matches, predictions)]
                        nmatch = sum(e <= tolerance_x for e in errors)
                        score = nmatch - sum(min(e / tolerance_x, 2) for e in errors) * .08
                        if best_x is None or score > best_x[0]:
                            positions = [round(x if e <= tolerance_x else p) for x, p, e in zip(matches, predictions, errors)]
                            best_x = score, positions, nmatch
            printed_rows = profile["printed_rows"]
            wanted = printed_rows + 2  # table top, header bottom, record boundaries
            tolerance_y = max(4, round(5 * h / 1988))
            best_y = None
            for top in y_candidates:
                if not .06 * h <= top <= .13 * h:
                    continue
                for step in np.arange(h * .0195, h * .0236, max(.15, h * .00012)):
                    bottom = top + (wanted - 1) * step
                    if not .88 * h <= bottom <= .99 * h:
                        continue
                    predictions = [top + i * step for i in range(wanted)]
                    matches = [min(y_candidates, key=lambda v: abs(v - p)) for p in predictions]
                    errors = [abs(y - p) for y, p in zip(matches, predictions)]
                    nmatch = sum(e <= tolerance_y for e in errors)
                    score = nmatch - sum(min(e / tolerance_y, 2) for e in errors) * .04
                    if best_y is None or score > best_y[0]:
                        positions = [round(y if e <= tolerance_y else p) for y, p, e in zip(matches, predictions, errors)]
                        best_y = score, positions, nmatch
            if best_x and best_y:
                x_positions, y_positions = best_x[1], best_y[1]
                if all(b > a for a, b in zip(x_positions, x_positions[1:])) and all(b > a for a, b in zip(y_positions, y_positions[1:])):
                    return x_positions, y_positions, (int(w), int(h)), {"method": "hough-template", "x_detected": int(best_x[2]), "x_total": len(x_positions), "y_detected": int(best_y[2]), "y_total": len(y_positions)}
        except ImportError:
            pass
    dark = im < 165
    ys = clusters(np.flatnonzero(dark[:, int(w * .07):int(w * .93)].mean(axis=1) > .59).tolist())
    xs = clusters(np.flatnonzero(dark[int(h * .09):int(h * .91), :].mean(axis=0) > .47).tolist())
    ys = [y for y in ys if int(h * .05) <= y <= int(h * .98)]
    xs = [x for x in xs if int(w * .03) <= x <= int(w * .97)]
    return xs, ys, (w, h), {"method": "projection", "x_detected": len(xs), "x_total": len(xs), "y_detected": len(ys), "y_total": len(ys)}


def candidates(args):
    workdir = Path(args.workdir).resolve()
    doc = manifest(workdir)
    profile = jread(Path(args.profile))
    overrides_file = workdir / "grid_overrides.json"
    overrides = jread(overrides_file) if overrides_file.exists() else {}
    columns = profile["columns"]
    allowed_x = profile.get("allowed_column_boundaries", [profile["expected_column_boundaries"]])
    all_rows, draft, issues = [], [], []
    for item in doc["pdfs"]:
        for page in selected_pages(item, args.limit_pages):
            image = page_path(workdir, item["stem"], page)
            ocr_file = ocr_path(workdir, item["stem"], page)
            if not image.is_file() or not ocr_file.is_file():
                issues.append({"pdf": item["name"], "page": page, "problem": "图片或OCR缺失"})
                continue
            xs, ys, size, quality = grid_lines(image, profile)
            override = overrides.get(item["name"], {}).get(str(page))
            if override:
                xs, ys = override["x"], override["y"]
                if any(b <= a for a, b in zip(xs, xs[1:])) or any(b <= a for a, b in zip(ys, ys[1:])):
                    raise ValueError(f"格线覆盖值未递增：{item['name']} 第{page}页")
                quality = {"method": "人工格线覆盖", "x_detected": len(xs), "x_total": len(xs), "y_detected": len(ys), "y_total": len(ys)}
            if len(xs) not in allowed_x or len(ys) != profile.get("printed_rows", len(ys) - 2) + 2:
                issues.append({"pdf": item["name"], "page": page, "problem": "格线数不符，须人工定位", "x_lines": xs, "y_lines": ys, "quality": quality})
                continue
            if quality["method"] not in ("hough-template", "人工格线覆盖") or quality["x_detected"] < len(xs) * .8 or quality["y_detected"] < len(ys) * .8:
                issues.append({"pdf": item["name"], "page": page, "problem": "部分格线按模板推算，须核对位置", "quality": quality})
            tokens = jread(ocr_file)["tokens"]
            rows = [{"pdf": item["name"], "page": page, "row": r, "image": str(image), "ocr_cells": {k: [] for k in columns}, "ocr_tokens": [], "grid": {"x": xs, "y_top": ys[r], "y_bottom": ys[r + 1], "size": size, "quality": quality}} for r in range(1, len(ys) - 1)]
            for token in tokens:
                poly = token.get("poly")
                if not poly or not isinstance(poly[0], (list, tuple)):
                    continue
                cx = sum(float(p[0]) for p in poly) / len(poly)
                cy = sum(float(p[1]) for p in poly) / len(poly)
                col = bisect.bisect_right(xs, cx) - 1
                yidx = bisect.bisect_right(ys, cy) - 1
                if 0 <= col < len(columns) and 1 <= yidx < len(ys) - 1:
                    record = rows[yidx - 1]
                    record["ocr_cells"][columns[col]].append(token["text"])
                    record["ocr_tokens"].append({**token, "column_candidate": columns[col]})
            for record in rows:
                record["ocr_cells"] = {k: " ".join(v).strip() for k, v in record["ocr_cells"].items()}
                all_rows.append(record)
                raw = record["ocr_cells"]
                draft.append({"pdf": record["pdf"], "page": page, "row": record["row"], "checked": False, "sample": None, "small": raw["小方"] or None, "tag": raw["标牌号"] or None, "species": raw["物种名"] or None, "x": raw["X"] or None, "y": raw["Y"] or None, "dbh": raw["DBH"] or None, **{f"d{i}": raw[f"D{i}"] or None for i in range(1, 7)}, "remark": raw["备注"] or None, "issue": "仅OCR候选，尚未看图复核"})
    write_jsonl(workdir / "candidates.jsonl", all_rows)
    write_jsonl(workdir / "reviewed_draft.jsonl", draft)
    jwrite(workdir / "grid_issues.json", issues)
    print(json.dumps({"candidate_rows": len(all_rows), "pages_needing_grid_review": len(issues)}, ensure_ascii=False))


def numeric(value, label, issues):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        issues.append(f"{label}非数值")
        return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    text = str(value).strip()
    if NUMBER.fullmatch(text):
        return float(text)
    issues.append(f"{label}原文“{text}”无法确认")
    return None


def branch(value, index, issues):
    if value is None or str(value).strip() == "":
        return None, None
    text = str(value).strip()
    prefix = BRANCH_PREFIX.fullmatch(text)
    if prefix:
        return float(prefix.group(2)), prefix.group(1).lower()
    match = BRANCH.fullmatch(text)
    if match:
        typ = match.group(2)
        suffix = match.group(3)
        if suffix and suffix != str(index):
            issues.append(f"D{index}后缀序号与列号不符：{text}")
        return float(match.group(1)), typ.lower() + (suffix or "") if typ else None
    issues.append(f"D{index}原文“{text}”无法拆分，正式列留空")
    return None, None


def remark_value(raw, profile, issues):
    if raw is None or not str(raw).strip():
        return None
    text = str(raw).strip()
    if profile.get("remark_mode") != "tree_status_only":
        return text
    keep, unknown = [], []
    for part in re.split(r"[；;,，]", text):
        part = part.strip()
        if not part:
            continue
        if any(re.search(pattern, part) for pattern in profile.get("remark_drop_patterns", [])):
            continue
        if any(re.search(pattern, part) for pattern in profile.get("remark_keep_patterns", [])):
            keep.append(part)
        else:
            unknown.append(part)
    if unknown:
        issues.append("备注内容未匹配树木状态规则，保留待核：" + "、".join(unknown))
    return "；".join(dict.fromkeys(keep + unknown)) or None


def names(args):
    workdir = Path(args.workdir).resolve()
    profile = jread(Path(args.profile))
    rows = list(jsonl(workdir / "reviewed.jsonl"))
    counts = collections.Counter(str(r.get("species") or "").strip() for r in rows)
    aliases = profile.get("species_aliases", {})
    audit = []
    for name, count in sorted(counts.items(), key=lambda x: (-x[1], x[0])):
        if not name:
            audit.append({"name": "", "count": count, "status": "原表空白，待核", "candidate": None, "source": None})
            continue
        alias = aliases.get(name)
        audit.append({"name": name, "count": count, "status": "有校正候选；仍须核对原图" if alias else "需对照原图和中文植物名录", "candidate": alias.get("name") if alias else None, "source": alias.get("source") if alias else None})
    jwrite(workdir / "species_audit.json", audit)
    print(json.dumps({"unique_names": len(audit), "requiring_image_review": sum(bool(x["name"]) for x in audit), "audit": str(workdir / "species_audit.json")}, ensure_ascii=False))


def finalize(args):
    workdir = Path(args.workdir).resolve()
    doc = manifest(workdir)
    profile = jread(Path(args.profile))
    page_reviews = jread(workdir / "page_reviews.json")
    reviewed = list(jsonl(workdir / "reviewed.jsonl"))
    by_location = {}
    for row in reviewed:
        missing = [key for key in REQUIRED if key not in row]
        if missing:
            raise ValueError(f"复核行缺少字段 {missing}: {row.get('pdf')} {row.get('page')}:{row.get('row')}")
        if row["checked"] is not True:
            raise ValueError(f"行尚未图像复核: {row['pdf']} {row['page']}:{row['row']}")
        key = (row["pdf"], int(row["page"]), int(row["row"]))
        if key in by_location:
            raise ValueError(f"重复页内位置: {key}")
        by_location[key] = row
    known = {x["name"] for x in doc["pdfs"]}
    if any(x[0] not in known for x in by_location):
        raise ValueError("reviewed.jsonl 包含 inventory 外的 PDF")
    page_count_by_name = {x["name"]: x["pages"] for x in doc["pdfs"]}
    if any(not 1 <= page <= page_count_by_name[pdf] for pdf, page, _ in by_location):
        raise ValueError("reviewed.jsonl 包含 PDF 页码范围外的记录")
    tag_pattern = re.compile(profile["tag_regex"])
    sample_pattern = re.compile(profile["sample_regex"])
    outputs = []
    for item in doc["pdfs"]:
        name = item["name"]
        page_spec = page_reviews.get(name, {})
        if {str(p) for p in range(1, item["pages"] + 1)} != set(page_spec):
            raise ValueError(f"{name} 页级复核不完整：应为1–{item['pages']}页")
        main, changes, summary, issues_by_loc = [], [], [], {}
        per_page = collections.defaultdict(list)
        for (pdf, page, rownum), record in by_location.items():
            if pdf == name:
                per_page[page].append(record)
        for page in range(1, item["pages"] + 1):
            review = page_spec[str(page)]
            if review.get("checked") is not True:
                raise ValueError(f"{name} 第{page}页未图像复核")
            page_rows = sorted(per_page[page], key=lambda r: int(r["row"]))
            expected = review.get("expected_rows")
            if not isinstance(expected, int) or expected < 0 or len(page_rows) != expected:
                raise ValueError(f"{name} 第{page}页：应有{expected}条，实有{len(page_rows)}条")
            if any(int(r["row"]) < 1 for r in page_rows):
                raise ValueError(f"{name} 第{page}页行号必须大于0")
            for r in page_rows:
                loc = (page, int(r["row"]))
                issue = [str(r["issue"])] if r.get("issue") else []
                sample = str(r["sample"]).strip() if r["sample"] is not None else ""
                tag = str(r["tag"]).strip() if r["tag"] is not None else ""
                if not sample_pattern.fullmatch(sample):
                    issue.append("样方编号缺失或格式不符")
                if not tag:
                    issue.append("标牌号空白")
                else:
                    m = tag_pattern.fullmatch(tag)
                    if not m:
                        issue.append("标牌号格式不符")
                    elif m.groupdict().get("sample") != sample:
                        issue.append("标牌号前缀与样方编号不一致")
                small = numeric(r["small"], "小方", issue)
                if small is None:
                    issue.append("小方空白")
                if small is not None and (not small.is_integer() or not profile["small_min"] <= small <= profile["small_max"]):
                    issue.append("小方超出配置范围")
                if r.get("small_source") and "?" in str(r["small_source"]):
                    issue.append("小方由问号格承接，需人工确认")
                species = str(r["species"]).strip() if r["species"] is not None else ""
                if re.search(r"[0-9?？]", species):
                    issue.append("物种名含数字或疑问号，需核对字形")
                alias = profile.get("species_aliases", {}).get(species)
                if alias:
                    issue.append(f"物种名可能需校正为{alias['name']}；须按原图字形确认")
                if not species:
                    issue.append("物种名空白")
                if not r.get("date"):
                    issue.append("调查日期空白")
                unit = r.get("xy_unit") or review.get("xy_unit")
                evidence = r.get("unit_evidence") or review.get("unit_evidence")
                x_raw = numeric(r["x"], "X", issue)
                y_raw = numeric(r["y"], "Y", issue)
                xy = []
                for axis, value in (("X", x_raw), ("Y", y_raw)):
                    if value is None:
                        issue.append(f"{axis}空白")
                        xy.append(None)
                    elif unit in ("cm", "m") and evidence:
                        converted = round(value / 100 if unit == "cm" else value, 6)
                        xy.append(converted)
                        if unit == "cm":
                            changes.append([name, page, loc[1], tag, axis, value, converted, "cm→m", evidence, "确定"])
                        if not profile["coordinate_min_m"] <= converted <= profile["coordinate_max_m"]:
                            issue.append(f"{axis}={converted} m 超出配置范围")
                    else:
                        xy.append(None)
                        issue.append(f"{axis}单位未由原图确认；原值{value}留在复核证据")
                dbh = numeric(r["dbh"], "DBH", issue)
                if dbh is None:
                    issue.append("DBH空白")
                elif dbh <= profile["dbh_min_exclusive"] or dbh > profile["dbh_soft_max_cm"]:
                    issue.append(f"DBH={dbh}超出配置范围")
                branches = []
                for i in range(1, 7):
                    diameter, typ = branch(r.get(f"d{i}"), i, issue)
                    branches.extend([diameter, typ])
                    if diameter is not None and diameter <= 0:
                        issue.append(f"D{i}={diameter}非正值")
                    if diameter is not None and dbh is not None and diameter > dbh:
                        issue.append(f"D{i}={diameter}大于DBH={dbh}")
                if r.get("status") == "待人工核对" and not issue:
                    issue.append("人工指定待核")
                remark = remark_value(r["remark"], profile, issue)
                if remark != r["remark"]:
                    changes.append([name, page, loc[1], tag, "备注", r["remark"], remark, "按树木状态规则整理", "原图已复核；其他原文保留于OCR/复核记录", "规则处理"])
                main.append([name, page, loc[1], r.get("date"), sample or None, int(small) if small is not None and small.is_integer() else small, tag or None, species or None, xy[0], xy[1], dbh, *branches, remark, "待人工核对" if issue else "图像复核通过", "；".join(dict.fromkeys(issue)) or None])
                issues_by_loc[loc] = bool(issue)
                for edit in r.get("edits", []):
                    if not all(k in edit for k in ("field", "from", "to", "evidence")):
                        raise ValueError(f"修订记录字段不完整: {name} {loc}")
                    changes.append([name, page, loc[1], tag, edit["field"], edit["from"], edit["to"], "图像/名录修订", edit["evidence"], r.get("confidence") or "未分级"])
            summary.append([name, page, expected, sum(issues_by_loc[(page, int(r["row"]))] for r in page_rows), review.get("xy_unit"), review.get("unit_evidence"), review.get("note")])
        tag_counts = collections.Counter(r[6] for r in main if r[6])
        duplicate_keys = {key for key, count in tag_counts.items() if count > 1}
        for row in main:
            if row[6] in duplicate_keys:
                row[-2] = "待人工核对"
                row[-1] = ((row[-1] + "；") if row[-1] else "") + "同一样方标牌号重复"
        pending_per_page = collections.Counter(r[1] for r in main if r[-2] == "待人工核对")
        for s in summary:
            s[3] = pending_per_page[s[1]]
        raw = []
        for page in range(1, item["pages"] + 1):
            source = ocr_path(workdir, item["stem"], page)
            if not source.exists():
                raise FileNotFoundError(f"OCR证据缺失: {source}")
            for idx, token in enumerate(jread(source)["tokens"], 1):
                raw.append([name, page, idx, token.get("text"), token.get("score"), json.dumps(token.get("poly"), ensure_ascii=False)])
        outputs.append({"source_pdf": item["path"], "source_sha256": item["sha256"], "file_name": item["stem"] + "_录入最终版.xlsx", "page_count": item["pages"], "main": main, "raw": raw, "summary": summary, "changes": changes, "record_count": len(main), "pending_count": sum(r[-2] == "待人工核对" for r in main)})
    payload = {"profile": profile["profile_name"], "outputs": outputs}
    jwrite(workdir / "final_payload.json", payload)
    print(json.dumps({"outputs": [{"file": x["file_name"], "rows": x["record_count"], "pending": x["pending_count"]} for x in outputs]}, ensure_ascii=False))


def validate(args):
    from openpyxl import load_workbook

    workdir = Path(args.workdir).resolve()
    payload = jread(workdir / "final_payload.json")
    output_dir = Path(args.output_dir).resolve()
    results = []
    for item in payload["outputs"]:
        source = Path(item["source_pdf"])
        if digest(source) != item["source_sha256"]:
            raise AssertionError(f"原 PDF 哈希变化：{source}")
        path = output_dir / item["file_name"]
        wb = load_workbook(path, read_only=False, data_only=True)
        if wb.sheetnames != ["主表", "原始OCR", "质检汇总", "修订记录"]:
            raise AssertionError((path, wb.sheetnames))
        main = wb["主表"]
        rows = list(main.values)
        if len(rows) - 1 != item["record_count"] or len(rows[0]) != 26:
            raise AssertionError(f"行列数不符：{path}")
        for number, (actual, expected) in enumerate(zip(rows[1:], item["main"]), 2):
            if list(actual) != expected:
                raise AssertionError(f"主表第{number}行与复核载荷不一致：{path}")
        if item["record_count"]:
            if len(main.tables) != 1 or not next(iter(main.tables.values())).autoFilter:
                raise AssertionError(f"主表筛选缺失：{path}")
        elif main.auto_filter.ref != "A1:Z1":
            raise AssertionError(f"空表筛选缺失：{path}")
        if sum(row[-2] == "待人工核对" for row in rows[1:]) != item["pending_count"]:
            raise AssertionError(f"待核条数不符：{path}")
        if len(wb["原始OCR"]["A"]) - 1 != len(item["raw"]):
            raise AssertionError(f"OCR条数不符：{path}")
        if len(wb["修订记录"]["A"]) - 1 != len(item["changes"]):
            raise AssertionError(f"修订日志条数不符：{path}")
        if any(not isinstance(row[7], str) and row[7] is not None for row in rows[1:]):
            raise AssertionError(f"物种名类型不符：{path}")
        results.append({"path": str(path), "pages": item["page_count"], "rows": item["record_count"], "pending": item["pending_count"], "changes": len(item["changes"])})
    print(json.dumps(results, ensure_ascii=False, indent=2))


def parser():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)
    d = sub.add_parser("doctor", help="检查本地依赖")
    d.add_argument("--input-dir")
    d.add_argument("--ocr-python", help="安装 PaddleOCR 的 Python 路径；也可设 PADDLE_OCR_PYTHON")
    d.set_defaults(func=doctor)
    i = sub.add_parser("inventory", help="只读列出PDF、页数与哈希")
    i.add_argument("--input-dir", required=True)
    i.add_argument("--workdir", required=True)
    i.add_argument("--exclude", action="append", default=[])
    i.set_defaults(func=inventory)
    for command, func in (("render", render), ("ocr", ocr), ("candidates", candidates), ("names", names), ("finalize", finalize)):
        a = sub.add_parser(command)
        a.add_argument("--workdir", required=True)
        if command in ("render", "ocr", "candidates"):
            a.add_argument("--limit-pages", type=int)
        if command == "render":
            a.add_argument("--dpi", type=int, default=220)
            a.add_argument("--force", action="store_true", help="重渲染并清除对应页的旧OCR候选")
        if command == "ocr":
            a.add_argument("--structure", action="store_true", help="改用 PP-StructureV3")
        if command in ("candidates", "names", "finalize"):
            a.add_argument("--profile", default=str(DEFAULT_PROFILE))
        a.set_defaults(func=func)
    v = sub.add_parser("validate", help="重新打开成品并核对")
    v.add_argument("--workdir", required=True)
    v.add_argument("--output-dir", required=True)
    v.set_defaults(func=validate)
    return p


if __name__ == "__main__":
    args = parser().parse_args()
    args.func(args)
